#!/usr/bin/env python3
"""
Gerador do dashboard de Análise de Variação (Orçado x Realizado) — FP&A.

Lê a base exportada do sistema (CSV), a classificação de contas e o de-para
oficial de segmentos, e produz um único arquivo HTML autossuficiente e
interativo, com os dados embutidos como JSON.

Uso:
    python gerar_dashboard.py

Saída:
    dashboard.html  (abra no navegador; funciona offline)

Todos os arquivos são procurados na mesma pasta deste script, então o programa
pode ser executado de qualquer diretório.
"""

import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime

try:
    import openpyxl
except ImportError:
    raise SystemExit("Dependência ausente. Instale com:  pip install openpyxl")


# ----------------------------------------------------------------------
# CONFIGURAÇÃO
# Caminhos absolutos, baseados na pasta onde este arquivo está.
# ----------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

XLSX_CLASSIF = os.path.join(BASE_DIR, "contas.xlsx")
XLSX_RESP = os.path.join(BASE_DIR, "responsaveis.xlsx")
CSV_PLANO = os.path.join(BASE_DIR, "plano.csv")
TEMPLATE = os.path.join(BASE_DIR, "template.html")
SAIDA = os.path.join(BASE_DIR, "dashboard.html")

# Nome das abas esperadas nas planilhas de apoio.
ABA_CLASSIF = "Planilha1"
ABA_RESP = "Relação"

# Codificação do CSV exportado do sistema.
CSV_ENCODING = "latin1"
CSV_DELIM = ";"

# Unidades que NÃO entram na análise de pendências nem nos segmentos, mas
# continuam sendo lidas: elas não deveriam ter lançamento nenhum, então
# qualquer movimento nelas é exceção e precisa ser repassado na hora.
UNIDADES_VIGIADAS = {
    "4801A",  # HUB - Maringá
    "4419A",  # CLIA Curitiba
    "5106A",  # Contratos Logísticos
    "4402A",  # CD Guarulhos
    "5110A",  # CD Mauá 1
    "5003A",  # Operação Portuária Salvador
    "5103A",  # Transportes Londrina
}

# Unidades que devem ser somadas em outra (origem -> destino).
UNIFICAR = {
    "4417A": "4418A",   # CD Químico Curitiba -> CD Curitiba
    "5002A": "5001A",   # HUB Salvador        -> Transportes Salvador
}

# Códigos que mudaram: a base financeira ainda usa o antigo, o cadastro já
# usa o novo (base -> cadastro).
COD_EQUIV = {
    "4001A": "4103A",   # Corporativo Itajaí
}

# Unidades ausentes do cadastro oficial: atribuição manual.
# Revisar quando o cadastro for atualizado.
SEG_MANUAL = {
    "4402A": "CDs",           # CD Guarulhos
    "4425A": "Fronteiras",    # PS Foz do Iguaçu 2
    "5003A": "Transportes",   # Operação Portuária Salvador
}

# Contas que são a mesma coisa contábil lançadas em códigos diferentes
# (origem -> destino). O orçado e o realizado são somados no destino.
FUNDIR_CONTAS = {
    "4.1.1.01.1.98": "4.1.1.01.1.99",   # Provisão PPR Colaboradores -> PPR Colaboradores
}

# Conta extra tratada como recorrente, além das marcadas no arquivo de
# classificação.
CONTAS_REC_EXTRA = {"4.1.1.02.1.09"}   # Etiquetagem

MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN",
         "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

# Posição das colunas de Planejado e Realizado no CSV (0-based).
PLAN_COLS = [8, 10, 12, 14, 16, 18, 20, 22, 24, 26, 28, 30]
REAL_COLS = [9, 11, 13, 15, 17, 19, 21, 23, 25, 27, 29, 31]

# Índices das colunas descritivas do CSV.
COL_CONTA = 0
COL_CONTA_DESC = 1
COL_UNIDADE = 2
COL_UNIDADE_DESC = 3

# Maior índice que o CSV precisa ter para ser considerado válido.
COL_MINIMA = max(REAL_COLS + PLAN_COLS +
                 [COL_CONTA, COL_CONTA_DESC, COL_UNIDADE, COL_UNIDADE_DESC]) + 1

# Marcadores substituídos no template pelo JSON dos dados.
MARCADOR_DADOS = r"/\*__DADOS__\*/.*?/\*__FIM__\*/"


# ----------------------------------------------------------------------
# VERIFICAÇÃO DOS ARQUIVOS DE ENTRADA
# ----------------------------------------------------------------------
def verifica_arquivos():
    """Confere se todos os arquivos necessários existem antes de processar."""
    obrigatorios = [
        (CSV_PLANO, "base exportada do sistema (planejado x realizado)"),
        (XLSX_CLASSIF, "classificação das contas (recorrente / fixo-variável)"),
        (XLSX_RESP, "de-para oficial de unidade para segmento"),
        (TEMPLATE, "modelo HTML do painel"),
    ]
    faltando = [(caminho, papel) for caminho, papel in obrigatorios
                if not os.path.exists(caminho)]
    if not faltando:
        return

    print("Não foi possível gerar o painel: arquivo(s) não encontrado(s).\n")
    for caminho, papel in faltando:
        print(f"  ausente: {os.path.basename(caminho)}")
        print(f"           {papel}")
    print(f"\nColoque o(s) arquivo(s) nesta pasta:\n  {BASE_DIR}")
    sys.exit(1)


def abre_planilha(caminho, aba, somente_leitura=False):
    """Abre uma aba de planilha com mensagem clara quando algo falta."""
    try:
        wb = openpyxl.load_workbook(caminho, data_only=True,
                                    read_only=somente_leitura)
    except Exception as erro:
        raise SystemExit(
            f"Não foi possível abrir {os.path.basename(caminho)}: {erro}")
    if aba not in wb.sheetnames:
        disponiveis = ", ".join(wb.sheetnames) or "(nenhuma)"
        wb.close()
        raise SystemExit(
            f"A planilha {os.path.basename(caminho)} não tem a aba '{aba}'.\n"
            f"Abas encontradas: {disponiveis}")
    return wb, wb[aba]


# ----------------------------------------------------------------------
# LEITURA E LIMPEZA
# ----------------------------------------------------------------------
def limpa(valor):
    """Remove o embrulho ="..." do Excel e aspas sobrando."""
    texto = (valor or "").strip()
    achado = re.match(r'^="?(.*?)"?$', texto)
    if achado:
        texto = achado.group(1)
    return texto.strip('"').strip()


def numero(valor):
    """Converte '1.234,56' (pt-BR) em float."""
    texto = limpa(valor).replace(".", "").replace(",", ".")
    try:
        return round(float(texto), 2)
    except ValueError:
        return 0.0


def celula(linha, indice):
    """Lê uma posição da linha do CSV sem estourar quando ela é mais curta."""
    if 0 <= indice < len(linha):
        return linha[indice]
    return ""


def carrega_classificacao(caminho):
    """De-para conta -> (recorrência, fixo/variável)."""
    wb, ws = abre_planilha(caminho, ABA_CLASSIF)
    clas = {}
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not linha or not linha[0]:
            continue
        conta = str(linha[0]).strip()
        recorrencia = linha[2] if len(linha) > 2 and linha[2] else "—"
        fixo_variavel = linha[3] if len(linha) > 3 and linha[3] else "—"
        clas[conta] = (recorrencia, fixo_variavel)
    wb.close()
    return clas


def carrega_segmentos(caminho):
    """De-para oficial unidade -> segmento (arquivo de responsáveis).

    O arquivo tem várias linhas por unidade (uma por centro de custo);
    fica o segmento mais frequente de cada uma.
    """
    wb, ws = abre_planilha(caminho, ABA_RESP, somente_leitura=True)
    votos = defaultdict(Counter)
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not linha or len(linha) < 6:
            continue
        unidade, segmento_lido = linha[0], linha[5]
        if unidade and segmento_lido:
            votos[str(unidade).strip()][str(segmento_lido).strip()] += 1
    wb.close()
    return {unidade: contagem.most_common(1)[0][0]
            for unidade, contagem in votos.items()}


def segmento(cod, mapa):
    """Segmento da unidade. Códigos 6xxx-9xxx são estruturas de Backoffice."""
    if cod[:1] in ("6", "7", "8", "9"):
        return "Backoffice"
    chave = COD_EQUIV.get(cod, cod)      # resolve código renomeado
    if chave in mapa:
        return mapa[chave]
    return SEG_MANUAL.get(cod, "Sem segmento")


def carrega_base(caminho_csv, clas, mapa_seg):
    """Lê o CSV, filtra a família 'A' e agrega por conta x unidade x mês."""
    try:
        with open(caminho_csv, encoding=CSV_ENCODING, newline="") as fh:
            linhas = list(csv.reader(fh, delimiter=CSV_DELIM))
    except UnicodeDecodeError:
        with open(caminho_csv, encoding="utf-8", newline="") as fh:
            linhas = list(csv.reader(fh, delimiter=CSV_DELIM))

    if len(linhas) < 2:
        raise SystemExit(
            f"{os.path.basename(caminho_csv)} não tem linhas de dados.")

    cabecalho = linhas[0]
    if len(cabecalho) < COL_MINIMA:
        raise SystemExit(
            f"{os.path.basename(caminho_csv)} tem {len(cabecalho)} colunas, "
            f"mas o esperado são pelo menos {COL_MINIMA}.\n"
            "Confirme se o export é o do plano com Planejado e Realizado "
            "de janeiro a dezembro.")

    comb = defaultdict(lambda: [[0.0] * 12, [0.0] * 12])   # (plan[12], real[12])
    unit_desc, acct_desc = {}, {}

    for linha in linhas[1:]:
        if not linha:
            continue
        unidade_original = limpa(celula(linha, COL_UNIDADE))
        if not unidade_original.endswith("A") or unidade_original == "DUMMYA":
            continue                       # usa só a família analítica "A"
        unidade = UNIFICAR.get(unidade_original, unidade_original)

        conta = limpa(celula(linha, COL_CONTA))
        if not conta:
            continue
        conta = FUNDIR_CONTAS.get(conta, conta)   # funde contas equivalentes

        if unidade not in unit_desc or unidade_original == unidade:
            unit_desc[unidade] = limpa(celula(linha, COL_UNIDADE_DESC))
        if conta not in acct_desc:
            acct_desc[conta] = limpa(celula(linha, COL_CONTA_DESC))

        planejado, realizado = comb[(conta, unidade)]
        for i in range(12):
            planejado[i] += numero(celula(linha, PLAN_COLS[i]))
            realizado[i] += numero(celula(linha, REAL_COLS[i]))

    if not comb:
        raise SystemExit(
            f"Nenhuma linha aproveitável em {os.path.basename(caminho_csv)}.\n"
            "Só entram unidades da família analítica (código terminado em 'A').")

    def compacta(valor):
        """Mantém os centavos, mas grava inteiro quando não há fração.

        Só reduz o tamanho do arquivo; o valor não muda.
        """
        arredondado = round(valor, 2)
        inteiro = int(arredondado)
        return inteiro if arredondado == inteiro else arredondado

    data = []
    for (conta, unidade), (planejado, realizado) in comb.items():
        tem_valor = (any(abs(x) > 0.005 for x in planejado)
                     or any(abs(x) > 0.005 for x in realizado))
        if tem_valor:
            data.append({"a": conta, "u": unidade,
                         "p": [compacta(x) for x in planejado],
                         "r": [compacta(x) for x in realizado]})

    units = [{"cod": unidade,
              "nome": unit_desc.get(unidade, unidade),
              "seg": ("Vigiada" if unidade in UNIDADES_VIGIADAS
                      else segmento(unidade, mapa_seg)),
              "vig": 1 if unidade in UNIDADES_VIGIADAS else 0}
             for unidade in sorted(unit_desc)]

    def eh_recorrente(cod):
        if cod in CONTAS_REC_EXTRA:
            return True
        recorrencia = str(clas.get(cod, ("", ""))[0]).lower()
        return recorrencia.startswith("recorrente")

    def natureza(cod):
        """Grupo contábil da conta.

        A receita não entra na conta de pendências: não faturar não é falha
        de lançamento. Ela serve para medir resultado.
        """
        prefixo = cod.split(".")[0]
        sub = cod[:3]
        if prefixo == "3":
            eh_deducao = cod[:5] == "3.1.1" and "(-)" in acct_desc.get(cod, "")
            return "ded" if eh_deducao else "rec"
        if prefixo == "4":
            if sub == "4.2":
                return "out"
            if sub == "4.3":
                return "fin"
            return "cd"
        if prefixo == "5":
            return "imp"
        return "out"

    accts = [{"cod": conta,
              "desc": acct_desc[conta],
              "nat": natureza(conta),
              "rec": clas.get(conta, ("—", "—"))[0],
              "fv": clas.get(conta, ("—", "—"))[1],
              "r": 1 if eh_recorrente(conta) else 0}
             for conta in sorted(acct_desc)]

    # Régua: um mês está "fechado" se tem qualquer realizado.
    mes_abs = [0.0] * 12
    for registro in data:
        for i in range(12):
            mes_abs[i] += abs(registro["r"][i])
    fechado = [abs(x) > 1 for x in mes_abs]

    return {"meses": MESES, "fechado": fechado,
            "units": units, "accts": accts, "data": data}


# ----------------------------------------------------------------------
# MONTAGEM DO HTML
# O front-end fica em template.html, ao lado deste arquivo. Separado de
# propósito: facilita editar HTML/CSS/JS com apoio do editor e deixa o
# histórico do Git legível.
# ----------------------------------------------------------------------
def carrega_template():
    with open(TEMPLATE, encoding="utf-8") as fh:
        return fh.read()


def monta_html(payload):
    """Injeta o JSON dos dados no template.

    No template o valor entre os marcadores é `null`, para que o arquivo
    continue sendo JavaScript válido caso alguém o abra direto.
    """
    payload["gerado"] = datetime.now().strftime("%Y-%m-%dT%H:%M")
    dados_json = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))

    html = carrega_template()
    # lambda no lugar do texto evita que "\\" e "\\g" do JSON sejam lidos
    # como referência de grupo pelo re.sub
    novo, trocas = re.subn(MARCADOR_DADOS, lambda _: dados_json,
                           html, count=1, flags=re.S)
    if trocas != 1:
        raise SystemExit(
            f"Marcador de dados não encontrado em {os.path.basename(TEMPLATE)}.\n"
            "O modelo precisa conter a linha:\n"
            "  const DB = /*__DADOS__*/ null /*__FIM__*/;")
    return novo


# ----------------------------------------------------------------------
def main():
    verifica_arquivos()

    clas = carrega_classificacao(XLSX_CLASSIF)
    mapa_seg = carrega_segmentos(XLSX_RESP)
    payload = carrega_base(CSV_PLANO, clas, mapa_seg)

    html = monta_html(payload)
    with open(SAIDA, "w", encoding="utf-8") as fh:
        fh.write(html)

    kb = os.path.getsize(SAIDA) / 1024
    print(f"OK -> {os.path.basename(SAIDA)} ({kb:.0f} KB)")
    print(f"   {len(payload['units'])} unidades, "
          f"{len(payload['data'])} combinações conta x unidade")

    fechados = [payload["meses"][i] for i in range(12) if payload["fechado"][i]]
    print(f"   meses fechados: {', '.join(fechados) if fechados else 'nenhum'}")

    for seg, qtd in sorted(Counter(u["seg"] for u in payload["units"]).items()):
        print(f"   {seg}: {qtd} unidades")

    n_rec = sum(a["r"] for a in payload["accts"])
    print(f"   contas recorrentes: {n_rec} de {len(payload['accts'])}")
    print("   contas por natureza:",
          dict(Counter(a["nat"] for a in payload["accts"])))

    vigiadas = [u for u in payload["units"] if u["vig"]]
    print(f"   unidades vigiadas: {len(vigiadas)}")

    orfas = [u for u in payload["units"] if u["seg"] == "Sem segmento"]
    if orfas:
        print("   ATENCAO - unidades sem segmento definido:")
        for unidade in orfas:
            print(f"      {unidade['cod']} {unidade['nome']}")


if __name__ == "__main__":
    main()
